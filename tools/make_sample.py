"""theme.json から動作確認用のサンプルデッキを生成する。

    python3 tools/make_sample.py            # examples/sample.pptx を出力
    python3 tools/make_sample.py out.pptx   # 出力先を指定

中身はダミーだが、カバー / セクション / KPI カード / 進捗バー / ステップフロー /
テーブル / コンテンツ / エンディングを一通り含み、theme.json の配色が反映される。
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from brandkit import Brand  # noqa: E402
from pptx.util import Inches  # noqa: E402


def main(out="examples/sample.pptx"):
    b = Brand()
    prs = b.new_presentation()
    b.delete_all_slides(prs)

    # 1. カバー
    b.add_cover(prs, "四半期サマリー", "2026.01.01")

    # 2. セクション
    b.add_section(prs, "01. ハイライト")

    # 3. KPI カード + 進捗バー（contentVisual）
    s = b.add_content(prs, "今四半期のハイライト", visual=True)
    margin = Inches(b.theme["style"]["pageMarginIn"])
    content_w = prs.slide_width - margin * 2
    b.add_kpi_row(s, [("128", "新規導入"), ("92%", "継続率"), ("4.6", "満足度")],
                  margin, Inches(2.0), content_w, Inches(1.6))
    b.add_progress_bar(s, margin, Inches(4.2), content_w, Inches(0.3), 0.72)

    # 4. ステップフロー（contentVisual）
    s = b.add_content(prs, "導入の流れ", visual=True)
    b.add_step_flow(s, ["申し込み", "初期設定", "本番運用"], margin, Inches(2.4), content_w)

    # 5. テーブル（contentVisual）
    s = b.add_content(prs, "プラン比較", visual=True)
    b.add_data_table(
        s, margin, Inches(2.0),
        ["プラン", "月額", "サポート"],
        [["Free", "0", "コミュニティ"], ["Pro", "1,980", "メール"], ["Team", "要相談", "専任"]],
        content_w,
    )

    # 6. テキスト中心コンテンツ
    s = b.add_content(prs, "次のアクション", visual=False)
    body = b.ph(s, 1)
    if body is not None:
        tf = body.text_frame
        tf.text = "オンボーディング資料の整備"
        for extra in ["フィードバック収集の仕組み化", "次四半期の目標設定"]:
            tf.add_paragraph().text = extra
        b.style_text(tf, b.TS["body"], b.C["dark"], font=b.F["bodyCJK"])

    # 7. エンディング
    b.add_ending(prs)

    out_path = Path(out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"saved: {out_path}  (slides={len(prs.slides._sldIdLst)}, "
          f"template={'custom' if b.using_custom else 'builtin-default'})")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "examples/sample.pptx")
